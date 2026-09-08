"""Importazione una tantum dei dati dal file Excel di partenza.

L'Excel e' solo la sorgente iniziale: dopo questa importazione l'app e'
autonoma e il file non viene piu' letto ne' scritto.
"""

import sqlite3
from datetime import date, datetime, time
from pathlib import Path

import openpyxl

from . import db


def _testo(valore) -> str | None:
    if valore is None:
        return None
    testo = str(valore).strip()
    return testo or None


def _data(valore) -> str | None:
    if isinstance(valore, datetime):
        return valore.date().isoformat()
    if isinstance(valore, date):
        return valore.isoformat()
    return None


def _ora(valore) -> str | None:
    if isinstance(valore, datetime):
        return valore.strftime("%H:%M")
    if isinstance(valore, time):
        return valore.strftime("%H:%M")
    return None


def _spezza_nome(intero: str) -> tuple[str, str]:
    """'Maria Rossi Bianchi' -> ('Maria', 'Rossi Bianchi')."""
    parti = intero.strip().split(None, 1)
    return (parti[0], parti[1]) if len(parti) == 2 else (intero.strip(), "")


def trova_persona(conn: sqlite3.Connection, nome_intero: str) -> int:
    """Restituisce l'id della persona, creandola se non esiste."""
    nome, cognome = _spezza_nome(nome_intero)
    riga = conn.execute(
        "SELECT id FROM persona WHERE nome = ? AND cognome = ?", (nome, cognome)
    ).fetchone()
    if riga:
        return riga["id"]
    cursore = conn.execute(
        "INSERT INTO persona (nome, cognome) VALUES (?, ?)", (nome, cognome)
    )
    return cursore.lastrowid


def importa(percorso_excel: Path | str, conn: sqlite3.Connection) -> dict:
    wb = openpyxl.load_workbook(percorso_excel, data_only=True)
    adesso = datetime.now().isoformat(timespec="seconds")

    # --- anagrafica tutor (foglio Dati, colonna F) ---
    dati = wb["Dati"]
    for riga in range(2, dati.max_row + 1):
        nome = _testo(dati.cell(riga, 6).value)
        if nome:
            trova_persona(conn, nome)

    # --- testata del piano (foglio Piano ISO) ---
    piano_iso = wb["Piano ISO"]
    nome_risorsa = _testo(piano_iso["C4"].value)
    reparto = _testo(piano_iso["H4"].value)
    mansione = _testo(piano_iso["M4"].value)
    responsabile = _testo(piano_iso["C5"].value)
    tutor_principale = _testo(piano_iso["H5"].value)
    data_inizio = _data(piano_iso["M5"].value)

    persona_risorsa = trova_persona(conn, nome_risorsa)
    risorsa_id = conn.execute(
        """
        INSERT INTO risorsa
            (persona_id, reparto, mansione, responsabile_id, tutor_principale_id,
             data_inizio, motivo)
        VALUES (?, ?, ?, ?, ?, ?, ?)
        """,
        (
            persona_risorsa,
            reparto,
            mansione,
            trova_persona(conn, responsabile) if responsabile else None,
            trova_persona(conn, tutor_principale) if tutor_principale else None,
            data_inizio,
            "Formazione",
        ),
    ).lastrowid

    piano_id = conn.execute(
        "INSERT INTO piano (risorsa_id, creato_il) VALUES (?, ?)", (risorsa_id, adesso)
    ).lastrowid

    # --- moduli: righe 9..29 del Piano ISO ---
    moduli_per_codice: dict[str, int] = {}
    ordine = 0
    for riga in range(9, 30):
        codice = _testo(piano_iso.cell(riga, 1).value)
        if not codice:
            continue
        ordine += 1
        area = _testo(piano_iso.cell(riga, 2).value) or ""
        titolo = _testo(piano_iso.cell(riga, 3).value) or ""
        applicabile = _testo(piano_iso.cell(riga, 4).value) or "SI"
        modalita = _testo(piano_iso.cell(riga, 5).value)
        tutor = _testo(piano_iso.cell(riga, 6).value)
        tutor_id = trova_persona(conn, tutor) if tutor else None

        conn.execute(
            """
            INSERT INTO modulo_catalogo
                (codice, area, titolo, modalita_default, tutor_referente_default_id, ordine)
            VALUES (?, ?, ?, ?, ?, ?)
            ON CONFLICT(codice) DO NOTHING
            """,
            (codice, area, titolo, modalita, tutor_id, ordine),
        )
        moduli_per_codice[codice] = conn.execute(
            """
            INSERT INTO piano_modulo
                (piano_id, codice, area, titolo, applicabile, modalita,
                 tutor_referente_id, ordine, entro_il, verifica_chiusura,
                 verifica_efficacia, data_verifica, esito)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                piano_id, codice, area, titolo, applicabile, modalita, tutor_id, ordine,
                _data(piano_iso.cell(riga, 13).value),
                _testo(piano_iso.cell(riga, 14).value),
                _testo(piano_iso.cell(riga, 15).value),
                _data(piano_iso.cell(riga, 16).value),
                _testo(piano_iso.cell(riga, 17).value),
            ),
        ).lastrowid

    # --- sessioni ---
    sessioni_foglio = wb["Sessioni"]
    n_sessioni = 0
    senza_modulo = []
    for riga in range(2, sessioni_foglio.max_row + 1):
        data_sessione = _data(sessioni_foglio.cell(riga, 1).value)
        if not data_sessione:
            continue
        ora_inizio = _ora(sessioni_foglio.cell(riga, 3).value)
        ora_fine = _ora(sessioni_foglio.cell(riga, 4).value)
        if not ora_inizio or not ora_fine:
            continue

        codice = _testo(sessioni_foglio.cell(riga, 6).value)
        dettaglio = _testo(sessioni_foglio.cell(riga, 8).value)
        tutor_cella = _testo(sessioni_foglio.cell(riga, 9).value)
        stato = _testo(sessioni_foglio.cell(riga, 10).value) or "Pianificata"
        esito = _testo(sessioni_foglio.cell(riga, 11).value)
        note = _testo(sessioni_foglio.cell(riga, 12).value)

        modulo_id = moduli_per_codice.get(codice) if codice else None
        if modulo_id is None:
            senza_modulo.append(f"{data_sessione} {dettaglio or ''}".strip())

        sessione_id = conn.execute(
            """
            INSERT INTO sessione
                (piano_id, piano_modulo_id, data, ora_inizio, ora_fine, dettaglio,
                 stato, esito_verifica, note, creata_il, modificata_il)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                piano_id, modulo_id, data_sessione, ora_inizio, ora_fine, dettaglio,
                stato, esito.upper() if esito else None, note, adesso, adesso,
            ),
        ).lastrowid
        n_sessioni += 1

        for nome_tutor in (tutor_cella or "").split(","):
            nome_tutor = nome_tutor.strip()
            if nome_tutor:
                conn.execute(
                    "INSERT OR IGNORE INTO sessione_tutor (sessione_id, persona_id) VALUES (?, ?)",
                    (sessione_id, trova_persona(conn, nome_tutor)),
                )

    conn.commit()
    return {
        "piano_id": piano_id,
        "risorsa": nome_risorsa,
        "moduli": len(moduli_per_codice),
        "sessioni": n_sessioni,
        "persone": conn.execute("SELECT COUNT(*) c FROM persona").fetchone()["c"],
        "sessioni_senza_modulo": senza_modulo,
    }
